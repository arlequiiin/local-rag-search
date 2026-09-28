"""Чтение документов (.docx, .md, .txt) в единый формат с плейсхолдерами изображений."""

import re
import shutil
import uuid
from dataclasses import dataclass, field
from pathlib import Path

SUPPORTED_EXTENSIONS = {".docx", ".md", ".txt"}

IMAGE_PLACEHOLDER = "[[image: {name}]]"
IMAGE_PLACEHOLDER_RE = re.compile(r"\[\[image:\s*([^\]]+?)\s*\]\]")
_MD_IMAGE_RE = re.compile(r"!\[[^\]]*\]\(([^)\s]+)(?:\s+\"[^\"]*\")?\)")
_SECTION_SEPARATOR_RE = re.compile(r"^\s*-{3,}\s*$", re.MULTILINE)


@dataclass
class Document:
    """Одна инструкция: заголовок, текст и изображения, на которые текст ссылается."""

    title: str
    text: str
    source: str
    images: list[str] = field(default_factory=list)


def parse_file(path: str | Path, images_dir: str | Path, split_sections: bool = False) -> list[Document]:
    """Разбирает файл в список документов.

    При ``split_sections=True`` файл режется на отдельные инструкции по строкам ``---``,
    а первая строка каждой секции становится её заголовком.
    Изображения сохраняются в ``images_dir``, в тексте на их месте остаётся ``[[image: имя]]``.
    """
    path = Path(path)
    images_dir = Path(images_dir)
    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Неподдерживаемый формат {ext!r}, ожидается один из {sorted(SUPPORTED_EXTENSIONS)}")

    prefix = uuid.uuid4().hex[:8]
    if ext == ".docx":
        text = _read_docx(path, images_dir, prefix)
    elif ext == ".md":
        text = _read_markdown(path, images_dir, prefix)
    else:
        text = path.read_text(encoding="utf-8")

    if not split_sections:
        return [_make_document(path.stem, text, path.name)]

    documents = []
    for section in _SECTION_SEPARATOR_RE.split(text):
        lines = section.strip().split("\n", 1)
        title = lines[0].strip().lstrip("#").strip()
        if not title:
            continue
        body = lines[1].strip() if len(lines) > 1 else ""
        documents.append(_make_document(title, body, path.name))
    return documents


def _make_document(title: str, text: str, source: str) -> Document:
    images = list(dict.fromkeys(IMAGE_PLACEHOLDER_RE.findall(text)))
    return Document(title=title, text=text.strip(), source=source, images=images)


def _read_docx(path: Path, images_dir: Path, prefix: str) -> str:
    from docx import Document as DocxDocument

    doc = DocxDocument(str(path))
    parts = []
    counter = 0

    for paragraph in doc.paragraphs:
        # Картинка в docx хранится внутри run как элемент graphic со ссылкой на файл
        for run in paragraph.runs:
            if "graphic" not in run._element.xml:
                continue
            blip = run._element.xpath(".//a:blip")
            if not blip:
                continue
            rel_id = blip[0].get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
            if rel_id not in doc.part.related_parts:
                continue
            image = doc.part.related_parts[rel_id]
            counter += 1
            name = f"{prefix}_{counter}{Path(image.partname).suffix}"
            images_dir.mkdir(parents=True, exist_ok=True)
            (images_dir / name).write_bytes(image.blob)
            parts.append(IMAGE_PLACEHOLDER.format(name=name))

        if paragraph.text.strip():
            parts.append(paragraph.text.strip())

    return "\n".join(parts)


def _read_markdown(path: Path, images_dir: Path, prefix: str) -> str:
    text = path.read_text(encoding="utf-8")
    counter = 0

    for match in _MD_IMAGE_RE.finditer(text):
        target = match.group(1)
        source = path.parent / target
        # Просто пропускаем ссылки в интернет или несуществующие файлы
        if target.startswith(("http://", "https://")) or not source.is_file():
            continue
        counter += 1
        name = f"{prefix}_{counter}{source.suffix}"
        images_dir.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(source, images_dir / name)
        text = text.replace(match.group(0), IMAGE_PLACEHOLDER.format(name=name))

    return text
