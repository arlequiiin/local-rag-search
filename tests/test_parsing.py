import io

import pytest
from conftest import png_bytes
from docx import Document as DocxDocument

from rag_search.parsing import parse_file


def test_txt(tmp_path):
    path = tmp_path / "Сброс пароля.txt"
    path.write_text("Откройте настройки.\nНажмите Сбросить.", encoding="utf-8")
    [doc] = parse_file(path, tmp_path / "images")
    assert doc.title == "Сброс пароля"
    assert doc.source == "Сброс пароля.txt"
    assert "Нажмите Сбросить" in doc.text
    assert doc.images == []


def test_markdown_local_images_are_copied(tmp_path):
    (tmp_path / "shot.png").write_bytes(png_bytes())
    path = tmp_path / "guide.md"
    path.write_text("Шаг 1 ![скрин](shot.png)\nВнешняя ![x](https://example.com/a.png)", encoding="utf-8")
    images_dir = tmp_path / "images"

    [doc] = parse_file(path, images_dir)

    assert len(doc.images) == 1
    assert (images_dir / doc.images[0]).read_bytes() == png_bytes()
    assert f"[[image: {doc.images[0]}]]" in doc.text
    assert "https://example.com/a.png" in doc.text


def test_split_sections(tmp_path):
    path = tmp_path / "many.md"
    path.write_text(
        "# Первая\nТекст первой --- с дефисами внутри\n\n---\n\n## Вторая\nТекст второй\n---\n\n",
        encoding="utf-8",
    )
    docs = parse_file(path, tmp_path / "images", split_sections=True)
    assert [d.title for d in docs] == ["Первая", "Вторая"]
    assert docs[0].text == "Текст первой --- с дефисами внутри"


def test_docx_text_and_images(tmp_path):
    source = DocxDocument()
    source.add_paragraph("Откройте меню.")
    source.add_picture(io.BytesIO(png_bytes()))
    source.add_paragraph("Нажмите Сохранить.")
    path = tmp_path / "Инструкция.docx"
    source.save(path)
    images_dir = tmp_path / "images"

    [doc] = parse_file(path, images_dir)

    assert doc.title == "Инструкция"
    lines = doc.text.split("\n")
    assert lines[0] == "Откройте меню."
    assert lines[1] == f"[[image: {doc.images[0]}]]"
    assert lines[2] == "Нажмите Сохранить."
    assert (images_dir / doc.images[0]).exists()


def test_unsupported_extension(tmp_path):
    path = tmp_path / "file.pdf"
    path.write_bytes(b"%PDF")
    with pytest.raises(ValueError, match="Неподдерживаемый формат"):
        parse_file(path, tmp_path)
