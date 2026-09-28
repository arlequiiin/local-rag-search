import io

import pytest
from conftest import png_bytes
from docx import Document as DocxDocument
from fastapi.testclient import TestClient

from rag_search.api import app
from rag_search.service import get_knowledge_base


@pytest.fixture
def client(kb):
    app.dependency_overrides[get_knowledge_base] = lambda: kb
    yield TestClient(app)
    app.dependency_overrides.clear()


def upload(client, name, content, **form):
    return client.post("/documents", files={"file": (name, content)}, data=form)


def test_full_document_lifecycle(client, kb):
    response = upload(
        client, "Печать.txt", "Если принтер не печатает чек, замените бумагу.".encode(), tags="касса, принтер"
    )
    assert response.status_code == 201
    [doc_id] = response.json()
    assert (kb.settings.docs_dir / "Печать.txt").exists()

    hits = client.post("/search", json={"query": "принтер не печатает"}).json()
    assert hits[0]["doc_id"] == doc_id

    answer = client.post("/ask", json={"query": "принтер не печатает"}).json()
    assert answer["answer"] == "Ответ по контексту"
    assert answer["best_title"] == "Печать"

    [doc] = client.get("/documents").json()
    assert doc["tags"] == ["касса", "принтер"]

    assert client.patch(f"/documents/{doc_id}", json={"active": False}).json()["active"] is False
    assert client.get("/documents", params={"active_only": True}).json() == []

    assert client.delete(f"/documents/{doc_id}").status_code == 204
    assert client.delete(f"/documents/{doc_id}").status_code == 404


def test_split_sections_upload(client):
    body = "Первая\nтекст один\n---\nВторая\nтекст два".encode()
    response = upload(client, "many.txt", body, split_sections="true")
    assert len(response.json()) == 2


def test_duplicate_filenames_do_not_overwrite(client, kb):
    upload(client, "a.txt", b"first text here")
    upload(client, "a.txt", b"second text here")
    assert sorted(p.name for p in kb.settings.docs_dir.iterdir()) == ["a.txt", "a_1.txt"]


def test_path_traversal_in_filename_is_neutralised(client, kb):
    upload(client, "../../evil.txt", b"some text")
    assert (kb.settings.docs_dir / "evil.txt").exists()


def test_unsupported_upload(client):
    assert upload(client, "doc.pdf", b"%PDF").status_code == 422


def test_images_are_served_and_removed_with_document(client, kb):
    source = DocxDocument()
    source.add_paragraph("Нажмите кнопку настроек.")
    source.add_picture(io.BytesIO(png_bytes()))
    buffer = io.BytesIO()
    source.save(buffer)

    [doc_id] = upload(client, "guide.docx", buffer.getvalue()).json()
    [image] = client.get("/documents").json()[0]["images"]
    assert client.get(f"/images/{image}").content == png_bytes()
    assert client.get("/images/..%2F..%2Fsecret").status_code == 404

    client.delete(f"/documents/{doc_id}")
    assert client.get(f"/images/{image}").status_code == 404


def test_health(client):
    assert client.get("/health").json() == {"status": "ok", "chunks": 0}
